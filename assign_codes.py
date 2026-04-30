#!/usr/bin/env python
# -*- coding: utf-8 -*-
import sys, re, os
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
from fuzzywuzzy import fuzz, process

INPUT_FILE   = 'ДС 9 Приложение №2 Укрупненная смета.xlsx'
OUTPUT_FILE  = 'ДС 9 Приложение №2 Укрупненная смета_READY.xlsx'
CLASSIFIER   = '2026-02-06_Классификатор статей СМР_как принимать.xlsx'
EXAMPLES_DIR = 'Пример смет'
SHEET_NAME   = 'Статьи СМР'

CODE_RE = re.compile(r"^(\d+\.(?:\d+\.)*)")

def extract_code(text):
    if not text: return ""
    m = CODE_RE.match(str(text).strip())
    if m: return m.group(1)
    raw = str(text).strip()
    if re.match(r"^\d+(\.\d+)*\.?$", raw):
        if not raw.endswith("."): raw += "."
        return raw
    return ""

def norm(text) -> str:
    return str(text).strip().lower() if text is not None else ""


def contains(name_n, *keywords):
    """Проверяет наличие любого из ключевых слов в нормализованном тексте."""
    return any(kw in name_n for kw in keywords)


def map_code(name_orig: str, num_b: str = "") -> str:
    """
    Возвращает кодификатор для строки по её наименованию.
    Возвращает пустую строку если не удалось определить код.
    """
    if not name_orig:
        return ""

    n = norm(name_orig)

    # ─── ПРИОРИТЕТ 1: Реставрация ────────────────────────────────────────────
    if contains(n, "реставрац", "воссоздани", "вычинка", "инъектировани",
                "инъецировани"):
        if contains(n, "фасад", "кладк", "лепнин"):
            return "19.1."
        if contains(n, "окно", "витраж", "двер"):
            return "19.2."
        if contains(n, "кровля"):
            return "19.3."
        if contains(n, "интерьер", "внутренн", "стена", "свод", "паркет",
                    "монолитн", "фундамент", "гидроизол", "молниезащ",
                    "подземн"):
            return "19.4."
        return "19.5."

    # ─── РАЗДЕЛ 16: Рабочая документация ─────────────────────────────────────
    if contains(n, "рабочая документация", "авторский надзор",
                 "рабочее проектировани", "архитектурные решения",
                 "конструктивные и объемно", "схема планировочн"):
        return "16."

    # ─── РАЗДЕЛ 18: Наружные сети ─────────────────────────────────────────────
    if contains(n, "наружные сети"):
        return "18."

    # ─── РАЗДЕЛ 1: Подготовительные работы ───────────────────────────────────
    if contains(n, "банковская гарантия", "обеспечение финансов"):
        return "1.1."
    if contains(n, "мобилизац", "демобилизац"):
        return "1.2."
    if contains(n, "бытовой городок", "строительный городок",
                 "временные офис", "временный офис", "ограждение строит",
                 "ограждение площадк", "ограждение стройплощадк",
                 "ограждение участк", "пункт мойки", "временные дороги",
                 "временных внутриплощадочн", "обустройство помещений",
                 "штаб строительства", "временная подсветка",
                 "перенос крун", "перенос ктпн"):
        return "1.2."
    if contains(n, "содержание площадки", "содержание временных",
                 "ежемесячная уборка", "уборка снега",
                 "меры предосторожности в зимн", "вывоз строительного мусора",
                 "ручные инструменты", "механизац", "кран башенн",
                 "коммунальн", "гарантийн", "охрана площадки",
                 "ежемесячная охрана", "генподрядн", "пропускная система",
                 "биометрический скуд", "обслуживание системы"):
        return "1.3."
    if contains(n, "итр ", "управление проектом", "вспомогательный персонал"):
        return "1.3."
    if contains(n, "геодезич"):
        return "1.3."
    if contains(n, "организация работ по проведению", "плановых проверок",
                 "согласование и совместный приказ", "открытие совместного приказ",
                 "техничес надзор"):
        return "1.3."
    if contains(n, "мокап"):
        return "1.4."
    if contains(n, "организация строительств", "этапность"):
        return "1.5."
    if contains(n, "ппр", "технологических карт", "лабораторн испытани"):
        return "1.3."
    if contains(n, "восстановление покрытия", "защита резиновым",
                 "защита металлическими", "содержание, ремонт и защита"):
        return "1.99."

    # ─── РАЗДЕЛ 2: Котлован ──────────────────────────────────────────────────
    if contains(n, "ограждение котлован", "стена в грунте", "свг",
                 "форшахта", "ограждающих конструкц котлован"):
        return "2.1."
    if contains(n, "срез оголовк", "срубка оголовк"):
        return "2.2."
    if contains(n, "распорная система", "шпунт", "извлечение шпунта"):
        return "2.3."
    if contains(n, "строительное водопонижени", "верхностный водоотлив",
                 "строительный водоотли"):
        return "2.5."
    if contains(n, "земляные работы", "разработка грунта", "вывоз грунта",
                 "перемещение грунта", "обратная засыпка", "устройство котлован"):
        return "2.6."
    if contains(n, "бнс", "буронабивн"):
        return "2.99."
    if contains(n, "котлован"):
        return "2."
    if contains(n, "демонтаж перекрытия в канале", "демонтаж бсс",
                 "демонтаж и переустройство подземных",
                 "демонтаж надземных сооружений тп",
                 "демонтаж конструкций"):
        return "1.99."
    if contains(n, "демонтаж", "ручная разборка", "ручной демонтаж"):
        return "1.2."
    if contains(n, "геологические изыскания"):
        return "1.99."
    if contains(n, "ручная разборка здания", "снос"):
        return "1.99."

    # ─── РАЗДЕЛ 3: Гидроизоляция подземной части ─────────────────────────────
    if contains(n, "виброизол", "виброзащитн"):
        return "3.3."
    if contains(n, "подготовка основания", "уплотнение основания",
                 "щебеночное основание"):
        return "3.1."
    if contains(n, "гидроизол"):
        if contains(n, "кровл"):
            return "7."
        if contains(n, "фундаментн плит"):
            return "3.2."
        if contains(n, "вертикальн"):
            return "3.2."
        if contains(n, "горизонтальн"):
            return "3.2."
        return "3.2."
    if contains(n, "дренаж"):
        return "3.99."

    # ─── РАЗДЕЛ 4: Конструктивные решения ────────────────────────────────────
    if contains(n, "металлоконструкц"):
        return "4.3."
    if contains(n, "усиление фундаментов", "усиление здания",
                 "усиление конструкций", "усиление вертикальн",
                 "усиление сводов", "устройство металлоконструкц в подземн",
                 "устройство ростверка", "устройство закладных"):
        return "4.99."
    if contains(n, "самонесущие ограждающ", "самонесущих конструкц"):
        return "4.99."
    if contains(n, "ж/б конструкции", "монолитн конструкц", "монолитные работы",
                 "возведение несущих", "несущих конструкций"):
        if contains(n, "подземн"):
            if contains(n, "горизонтальн", "перекрыти", "плит"):
                return "4.1.2."
            if contains(n, "вертикальн", "стен", "колонн", "пилон"):
                return "4.1.3."
            return "4.1."
        if contains(n, "надземн", "типовой этаж"):
            if contains(n, "горизонтальн", "перекрыти"):
                return "4.2.3."
            if contains(n, "вертикальн", "стен", "колонн"):
                return "4.2.4."
            return "4.2."
        return "4."
    if contains(n, "фундаментная плита", "фундаментной плит"):
        return "4.1.1."
    if contains(n, "устройство траншеи", "выемка грунта",
                 "установка стоек переопирани", "устройство фундаментной"):
        return "4.99."

    # ─── РАЗДЕЛ 5: Общестроительные работы ───────────────────────────────────
    if contains(n, "светопрозрачные перегородки"):
        return "5.3."
    if contains(n, "перегородки"):
        if contains(n, "квартир", "mr base"):
            return "14.1."
        if contains(n, "паркинг", "подземн"):
            return "5.1."
        return "5.2."

    # ─── РАЗДЕЛ 6: Фасадные работы ───────────────────────────────────────────
    if contains(n, "баннерная конструкция", "удерживающих конструкц",
                 "ипрс", "фасадной сетки", "самонесущая баннерн"):
        return "6.99."
    if contains(n, "разборка/восстановление исторических фасадов",
                 "разборка исторических фасадов"):
        return "19.1."
    if contains(n, "фасад"):
        if contains(n, "витраж", "стоечно-ригельная"):
            if contains(n, "стилобат", "1й", "1-й"):
                return "6.1."
            return "6.2."
        if contains(n, "навесн", "вентфасад", "нвф"):
            if contains(n, "стилобат", "1й", "цоколь"):
                return "6.3."
            return "6.4."
        if contains(n, "модульн"):
            return "6.5."
        if contains(n, "обслуживания фасадов", "соф"):
            return "6.6."
        if contains(n, "штукатурн"):
            return "6.7."
        if contains(n, "демонтаж фасадн"):
            return "6.8."
        if contains(n, "откос", "отлив", "парапет"):
            return "6.10."
        return "6."
    if contains(n, "витраж", "стоечно-ригельная"):
        if contains(n, "стилобат", "1й", "1-й"):
            return "6.1."
        return "6.2."
    if contains(n, "устройство фасадов"):
        return "6."

    # ─── РАЗДЕЛ 7: Кровля ─────────────────────────────────────────────────────
    if contains(n, "кровля", "кровли", "временная кровля"):
        if contains(n, "паркинг", "подземн", "стилобат"):
            return "7.1."
        if contains(n, "парапет", "лестниц", "ограждени"):
            return "7.3."
        if contains(n, "воронк", "люк"):
            return "7.99."
        return "7.2."

    # ─── РАЗДЕЛ 8: Отделочные работы МОП и паркинга ──────────────────────────
    if contains(n, "двери", "ворота", "противопожарная штора"):
        return "8.3."
    if contains(n, "навигация и обеспечение доступа"):
        return "13.99."
    if contains(n, "отделк"):
        if contains(n, "квартир", "mr base"):
            return "14."
        if contains(n, "паркинг", "подземн", "рамп"):
            return "8.1.1."
        if contains(n, "лобби", "гранд"):
            return "8.2.1."
        if contains(n, "воинской части"):
            return "8.2."
        return "8.2."

    # ─── РАЗДЕЛ 9: Лифты ─────────────────────────────────────────────────────
    if contains(n, "лифт", "подъёмник", "подъемник"):
        if contains(n, "отделка кабин"):
            return "9.3."
        if contains(n, "управление лифтами", "вызывная"):
            return "9.99."
        return "9.2."

    # ─── РАЗДЕЛ 10: ВИС механические ─────────────────────────────────────────
    if contains(n, "водоснабжени", "хвс", "гвс", "горячее водоснабж",
                 "холодное водоснабж", "хоз-питьевого водоснабж"):
        return "10.1."
    if contains(n, "газовое пожаротушени", "порошковое пожаротушени"):
        return "10.9."
    if contains(n, "водяное пожаротушени", "спринклерн",
                 "противопожарный водопровод", "аупт",
                 "автоматических установок пожаротушени"):
        return "10.2."
    if contains(n, "канализац", "водоотведени", "жироуловит",
                 "дренажных насосов"):
        return "10.3."
    if contains(n, "отоплени", "теплоснабж", "радиатор", "конвектор"):
        return "10.4."
    if contains(n, "противодымная вентиляц", "дымоудалени",
                 "подпор воздуха", "противодымн вентиляц"):
        return "10.7."
    if contains(n, "тепловой пункт", " итп", "итп "):
        return "10.8."
    if contains(n, "вентиляц"):
        if contains(n, "приточно-вытяжная") and contains(n, "квартир"):
            return "14.9."
        return "10.5."
    if contains(n, "кондиционировани"):
        if contains(n, "квартир"):
            return "14.10."
        return "10.6."

    # ─── РАЗДЕЛ 11: ВИС электрические и слаботочные ──────────────────────────
    if contains(n, "система электроснабжения и электроосвещения",
                 "внутреннее электрооборудование",
                 "электроснабжени", "электроосвещени",
                 "электромонтажн", "переустройство системы электроснабж"):
        if contains(n, "кабелен", "лотки"):
            return "11.1.1."
        if contains(n, "щитовое", "вру", "грщ", "щит", "распределительные силовые"):
            return "11.1.3."
        if contains(n, "шинопровод"):
            return "11.1.5."
        if contains(n, "зарядные устройства", "дгу", "ибп", "резервные источники"):
            return "11.1.8."
        return "11.1."
    if contains(n, "молниезащит", "заземлени",
                 "системы уравнивания потенциалов"):
        return "11.2."
    if contains(n, "сети связи"):
        return "11.3."
    if contains(n, "радиофикац", "го и чс", "го чс"):
        return "11.3.1."
    if contains(n, "структурированная кабельная система систем связи", "скс сс"):
        return "11.3.2."
    if contains(n, "стс мгн", "сигнализации для санузлов мгн"):
        return "11.3.3."
    if contains(n, "сусс", "усиления сотовой"):
        return "11.3.4."
    if contains(n, "wi-fi", "wifi", "беспроводная локальная вычилит"):
        return "11.3.5."
    if contains(n, "музыкальная трансляция"):
        return "11.3."
    if contains(n, "пожарная сигнализац", "апс", "соуэ", "аспс", "аспа",
                 "система пожарной сигнализации и автоматики"):
        return "11.4."
    if contains(n, "система оповещения и управления эвакуацией"):
        return "11.4."
    if contains(n, "противопожарные системы"):
        return "11.4."
    if contains(n, "скуд", "контроль и управление доступом"):
        return "11.5.3."
    if contains(n, "охранного видеонаблюдени", "сов ", "система охранного"):
        return "11.5.2."
    if contains(n, "сотс", "охранно-тревожной"):
        return "11.5.4."
    if contains(n, "скс сб", "систем безопасности"):
        return "11.5.1."
    if contains(n, "домофон", "сдс"):
        return "11.5."
    if contains(n, "система контроля и управления доступом",
                 "охранная сигнализац", "система охранного видеонаблюд"):
        return "11.5."
    if contains(n, "оздс", "дератизац"):
        return "11.10."
    if contains(n, "наружное освещени", "опоры освещени"):
        return "11.11."
    if contains(n, "фасадное освещени"):
        return "11.12."
    if contains(n, "есцз", "ецсз", "цифровой среды здания",
                 "цифровая среда здания"):
        return "11.8."
    if contains(n, "умная квартира", "умный офис"):
        return "11.9."
    if contains(n, "аскуэ", "учёта электроэнергии", "учета электроэнергии",
                 "учёта воды", "учета воды", "учёта тепла", "учета тепла",
                 "автоматизированная система учета", "аскувт", "асувт", "асуэ"):
        return "11.7."
    if contains(n, "автоматизац", "диспетчеризац", "диспетчерского контроля",
                 "асуд"):
        if contains(n, "контроль загазованност"):
            return "11.6.7."
        if contains(n, "лифт"):
            return "11.6.3."
        if contains(n, "противодымн"):
            return "11.6.2."
        if contains(n, "канализац", "дренаж"):
            return "11.6.4."
        if contains(n, "водоснабж"):
            return "11.6.6."
        if contains(n, "холодоснабж"):
            return "11.6.8."
        if contains(n, "пожаротушени"):
            return "11.6.5."
        if contains(n, "общеобменн вентиляц", "вентиляции"):
            return "11.6.1."
        return "11.6."

    # ─── РАЗДЕЛ 12: Благоустройство ──────────────────────────────────────────
    if contains(n, "благоустройств"):
        return "12."
    if contains(n, "тротуар", "брусчатка", "бортовой камень", "асфальт"):
        return "12.2."
    if contains(n, "озеленени", "деревья", "газон"):
        return "12.3."
    if contains(n, "маф", "детские площадки", "велопарковка", "фонтан"):
        return "12.4."

    # ─── РАЗДЕЛ 13: ТХ ───────────────────────────────────────────────────────
    if contains(n, "технологические решения"):
        if contains(n, "паркинг", "автостоянки"):
            return "13.1."
        if contains(n, "апс", "автоматизированная парковочная"):
            return "13.2."
        if contains(n, "мусоро"):
            return "13.3."
        return "13.4."
    if contains(n, "мусороудалени", "мусоропровод"):
        return "13.3."
    if contains(n, "система «анти", "противодействи террористич"):
        return "13.99."
    if contains(n, "навигация"):
        return "13.99."

    # ─── РАЗДЕЛ 14: MR Base ──────────────────────────────────────────────────
    if contains(n, "mr base", "mrbase"):
        return "14."
    if contains(n, "умная"):
        return "11.9."

    return ""


def load_classifier(path):
    mapping = {}
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[SHEET_NAME]
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r or len(r) < 11: continue
        kod = r[10]
        if not kod: continue
        kod = str(kod).strip()
        for name in [r[2], r[4], r[7]]:
            if name:
                key = norm(name)
                if key and key not in mapping:
                    mapping[key] = kod
    wb.close()
    return mapping

def load_examples(folder):
    mapping = {}
    for fname in sorted(os.listdir(folder)):
        if not fname.endswith(".xlsx"): continue
        fpath = os.path.join(folder, fname)
        try:
            wb = openpyxl.load_workbook(fpath, read_only=True)
        except Exception as e:
            print("  [WARN] %s: %s" % (fname, e))
            continue
        ws = wb.active
        header_row_num = 1
        code_cols, name_cols = [], []
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
            has_code = any(row[j] and 'кодификатор' in norm(row[j]) for j in range(len(row)))
            if has_code:
                header_row_num = i + 1
                for j, cell in enumerate(row):
                    if cell and 'кодификатор' in norm(cell): code_cols.append(j)
                    if cell and 'наименование' in norm(cell): name_cols.append(j)
                break
        for row in ws.iter_rows(min_row=header_row_num + 1, values_only=True):
            if not row: continue
            pairs = []
            if len(row) >= 7:
                if row[1] and row[2]: pairs.append((str(row[1]), str(row[2])))
                if row[3] and row[4]: pairs.append((str(row[3]), str(row[4])))
                if row[5] and row[6]: pairs.append((str(row[5]), str(row[6])))
                if len(row) >= 8 and row[7] and row[4]: pairs.append((str(row[7]), str(row[4])))
            if len(row) >= 4 and row[2] and row[3]: pairs.append((str(row[2]), str(row[3])))
            for ci, ni in zip(code_cols, name_cols):
                if ci < len(row) and ni < len(row) and row[ci] and row[ni]:
                    pairs.append((str(row[ci]), str(row[ni])))
            for raw_code, raw_name in pairs:
                code = extract_code(raw_code)
                if code and raw_name:
                    key = norm(raw_name)
                    if key and key not in mapping:
                        mapping[key] = code
        wb.close()
    print("  Examples: %d pairs loaded" % len(mapping))
    return mapping

def load_ds9(path):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        rows.append({
            "sheet_row": i,
            "num_b": str(row[1].value).strip() if row[1].value is not None else None,
            "code_c": str(row[2].value).strip() if row[2].value is not None else None,
            "name_d": row[3].value,
            "cell_c": row[2],
        })
    return wb, ws, rows

def build_hierarchy(rows):
    parent = {}
    for i, row in enumerate(rows):
        nb = row["num_b"]
        if not nb:
            parent[i] = None
            continue
        parts = [p for p in nb.split(".") if p]
        found = None
        for j in range(i - 1, -1, -1):
            pnb = rows[j]["num_b"]
            if not pnb: continue
            pparts = [p for p in pnb.split(".") if p]
            if len(pparts) < len(parts) and parts[:len(pparts)] == pparts:
                found = j
                break
        parent[i] = found
    return parent

INHERIT_KEYWORDS = [
    'окн1',
    'окн2',
    'окн-1',
    'окн-2',
    'корпус r',
    'основное оборудование',
    'кабельные линии',
    'оконечное оборудование',
    'прочее оборудование',
    'прочие работы',
    'пусконаладочные работы',
    'подготовительные работы',
    'завершающие работы',
    'пнр',
]

def is_generic(name):
    n = norm(name)
    return any(k in n for k in INHERIT_KEYWORDS)

def main():
    print("=" * 60)
    print('assign_codes.py — проставление кодов Статья СМР')
    print("=" * 60)
    print('\n[1] Загружаем классификатор...')
    cmap = load_classifier(CLASSIFIER)
    print("    Классификатор: %d записей" % len(cmap))
    print('\n[2] Загружаем примеры смет...')
    emap = load_examples(EXAMPLES_DIR)
    combined_map = {}
    combined_map.update(cmap)
    combined_map.update(emap)
    combined_keys = list(combined_map.keys())
    print("    Итого уникальных пар: %d" % len(combined_map))
    print('\n[3] Читаем ДС9 и строим иерархию...')
    wb, ws, rows = load_ds9(INPUT_FILE)
    parent_idx = build_hierarchy(rows)

    def get_ancestor_code(idx):
        visited = set()
        cur = parent_idx.get(idx)
        while cur is not None and cur not in visited:
            visited.add(cur)
            code = rows[cur].get("code_c") or ""
            if code: return code
            cur = parent_idx.get(cur)
        return ""

    print('\n[4] Проставляем коды...')
    filled = 0
    already = 0
    manual = []

    for i, row in enumerate(rows):
        code = row["code_c"]
        if code:
            already += 1
            continue
        name  = row["name_d"]
        num_b = row["num_b"]
        if not name and not num_b: continue
        name_n = norm(name) if name else ""
        assigned = ""
        if name and is_generic(str(name)):
            assigned = get_ancestor_code(i)
        if not assigned and name_n in combined_map:
            assigned = combined_map[name_n]
        if not assigned and name_n and len(name_n) >= 5:
            match = process.extractOne(
                name_n, combined_keys,
                scorer=fuzz.token_sort_ratio,
                score_cutoff=80
            )
            if match:
                assigned = combined_map[match[0]]
        if not assigned and name:
            assigned = map_code(str(name), str(num_b) if num_b else "")
        if not assigned and num_b and name:
            assigned = get_ancestor_code(i)
        if assigned:
            row["code_c"] = assigned
            row["cell_c"].value = assigned
            filled += 1
        else:
            manual.append("  строка %4d | B=%s | D=%s" % (
                row["sheet_row"], num_b, str(name)[:70] if name else ""))

    print('\n[5] Результаты:')
    print("    Уже были с кодом:     %d" % already)
    print("    Проставлено скриптом: %d" % filled)
    print("    Осталось пустых:      %d" % len(manual))
    print('\n[6] Сохраняем в ' + OUTPUT_FILE + "...")
    wb.save(OUTPUT_FILE)
    print("    Готово!")
    if manual:
        sep = "=" * 60
        print(""+sep)
        print('ТРЕБУЮТ РУЧНОЙ ПРОВЕРКИ (%d строк):' % len(manual))
        print(sep)
        for line in manual:
            print(line)
    print('\nФинально готово.')

if __name__ == "__main__":
    main()
