#!/usr/bin/env python3
"""
Сверка сметы генподрядчика с классификатором СМР.

Использование:
    python checker.py <классификатор.xlsx> <смета.xlsx> [-o отчёт.xlsx] [-t 85]
"""

import sys
import re
import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from fuzzywuzzy import fuzz

# ── Визуальные константы ───────────────────────────────────────────────────────
FILL_HEADER = PatternFill("solid", fgColor="4472C4")
FILL_OK     = PatternFill("solid", fgColor="C6EFCE")
FILL_WARN   = PatternFill("solid", fgColor="FFEB9C")
FILL_ERR    = PatternFill("solid", fgColor="FFC7CE")
FONT_HEADER = Font(bold=True, color="FFFFFF")

# Паттерн допустимых кодов: X  /  X.Y  /  X.Y.Z
CODE_RE = re.compile(r'^\d+(\.\d+){0,2}$')


# ── Вспомогательные функции ────────────────────────────────────────────────────

def norm(text) -> str:
    """Нижний регистр, обрезка пробелов."""
    return str(text).strip().lower() if text is not None else ""


def norm_unit(text) -> str:
    """Нормализация единиц измерения к каноническому виду."""
    u = norm(text).rstrip(".")
    mapping = {
        "кв.м": "м²", "кв. м": "м²", "м2": "м²",
        "куб.м": "м³", "куб. м": "м³", "м3": "м³",
        "пог.м": "п.м", "пог. м": "п.м", "п/м": "п.м", "пм": "п.м",
        "чел/ч": "чел.-ч", "чел.ч": "чел.-ч", "чел.-час": "чел.-ч",
        "тн": "т", "тонн": "т", "тонна": "т",
        "шт.": "шт", "компл.": "компл",
    }
    return mapping.get(u, u)


def is_code(val) -> bool:
    return bool(CODE_RE.match(str(val).strip())) if val is not None else False


def code_level(code: str) -> int:
    return len(code.strip().split("."))


def find_col(headers: tuple, keywords: list) -> int | None:
    """Найти индекс столбца по ключевым словам в заголовке (регистронезависимо)."""
    for i, h in enumerate(headers):
        hn = norm(h)
        for kw in keywords:
            if kw in hn:
                return i
    return None


def auto_width(ws):
    """Автоматическая ширина столбцов (не шире 55 символов)."""
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 55)


def write_header(ws, headers: list):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 36


# ── Загрузка классификатора ────────────────────────────────────────────────────

def load_classifier(path: str) -> dict:
    """
    Возвращает dict {code: {"name": str, "name_orig": str, "unit": str, "unit_orig": str}}.

    Поддерживает два формата Excel:
      1. Отдельные столбцы «Раздел» / «Статья» / «Позиция» + «Наименование» + «Ед.»
      2. Единый столбец с кодом (X / X.Y / X.Y.Z) + «Наименование» + «Ед.»
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    # Ищем строку заголовка
    header_idx = None
    for i, row in enumerate(rows):
        if any("наименование" in norm(c) for c in row if c is not None):
            header_idx = i
            break

    if header_idx is None:
        # Fallback: первая строка, где есть код
        for i, row in enumerate(rows):
            if any(is_code(c) for c in row):
                header_idx = max(i - 1, 0)
                break

    headers = rows[header_idx] if header_idx is not None else ()

    # Определяем столбцы
    раздел_col  = find_col(headers, ["раздел"])
    статья_col  = find_col(headers, ["статья"])
    позиция_col = find_col(headers, ["позиция", "пункт"])
    single_col  = find_col(headers, ["код", "кодификатор"])
    name_col    = find_col(headers, ["наименование"])
    unit_col    = find_col(headers, ["ед.", "ед ", "единиц"])

    # Если code-столбец не найден по заголовку — ищем по содержимому
    if single_col is None and раздел_col is None:
        start = (header_idx + 1) if header_idx is not None else 0
        for row in rows[start: start + 10]:
            for j, cell in enumerate(row):
                if is_code(cell):
                    single_col = j
                    break
            if single_col is not None:
                break

    # Позиционные fallback'и
    if name_col is None and single_col is not None:
        name_col = single_col + 1
    if unit_col is None and name_col is not None:
        unit_col = name_col + 1

    def safe(row, idx):
        return row[idx] if idx is not None and idx < len(row) else None

    classifier: dict = {}
    start = (header_idx + 1) if header_idx is not None else 0

    for row in rows[start:]:
        if all(c is None for c in row):
            continue

        # Определяем код строки: сначала проверяем позицию (X.Y.Z), затем статью (X.Y), затем раздел (X)
        code = None
        for col in (позиция_col, статья_col, раздел_col, single_col):
            if col is not None and is_code(safe(row, col)):
                code = str(safe(row, col)).strip()
                break

        if code is None:
            continue

        name_orig = str(safe(row, name_col)).strip() if safe(row, name_col) else ""
        unit_orig = str(safe(row, unit_col)).strip() if safe(row, unit_col) else ""

        classifier[code] = {
            "name":      norm(name_orig),
            "name_orig": name_orig,
            "unit":      norm_unit(unit_orig),
            "unit_orig": unit_orig,
        }

    return classifier


# ── Загрузка сметы ─────────────────────────────────────────────────────────────

def load_estimate(path: str) -> list:
    """
    Возвращает список строк сметы:
      [{"row_num", "code", "name", "name_orig", "unit", "unit_orig"}, ...]
    """
    wb = openpyxl.load_workbook(path, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        header_idx = None
        code_col = name_col = unit_col = None

        for i, row in enumerate(rows):
            rn = [norm(c) for c in row]
            if any("кодификатор" in c for c in rn) or any("наименование" in c for c in rn):
                header_idx = i
                code_col = find_col(row, ["кодификатор", "код"])
                name_col = find_col(row, ["наименование"])
                unit_col = find_col(row, ["ед.", "ед ", "единиц"])
                break

        if header_idx is None:
            continue

        def safe(row, idx):
            return row[idx] if idx is not None and idx < len(row) else None

        items = []
        for abs_row, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if all(c is None for c in row):
                continue
            code_val = safe(row, code_col)
            if not is_code(code_val):
                continue
            name_val = safe(row, name_col)
            unit_val = safe(row, unit_col)
            items.append({
                "row_num":   abs_row,
                "code":      str(code_val).strip(),
                "name":      norm(name_val),
                "name_orig": str(name_val).strip() if name_val else "",
                "unit":      norm_unit(unit_val),
                "unit_orig": str(unit_val).strip() if unit_val else "",
            })

        if items:
            return items

    return []


# ── Проверка строк ─────────────────────────────────────────────────────────────

def check(classifier: dict, estimate: list, threshold: int = 85) -> list:
    """
    Проверяет каждую строку сметы по трём критериям:
      - существование кода в классификаторе
      - соответствие наименования (нечёткое сравнение, порог = threshold %)
      - соответствие единиц измерения
    """
    results = []

    for item in estimate:
        code = item["code"]
        r = {
            "row_num":    item["row_num"],
            "code":       code,
            "name_est":   item["name_orig"],
            "unit_est":   item["unit_orig"],
            "status":     "ОК",
            "error_type": "",
            "expected":   "",
            "recommendation": "",
        }

        errors = []

        if code not in classifier:
            children = [k for k in classifier if k.startswith(code + ".")]
            if children:
                r["status"] = "УТОЧНЕНИЕ"
                errors.append("Требуется уточнение детализации")
                examples = ", ".join(children[:3])
                r["expected"] = f"Коды детализации: {examples}"
                r["recommendation"] = f"Уточните код, например: {examples[:60]}"
            else:
                r["status"] = "ОШИБКА"
                errors.append("Код не найден")
                r["recommendation"] = "Проверьте правильность кода кодификатора в смете"
        else:
            children = [k for k in classifier if k.startswith(code + ".")]
            if children:
                r["status"] = "УТОЧНЕНИЕ"
                errors.append("Требуется уточнение детализации")
                examples = ", ".join(children[:3])
                r["expected"] = f"Коды детализации: {examples}"
                r["recommendation"] = f"Уточните код до позиции, например: {examples[:60]}"
            else:
                clf = classifier[code]

                sim = fuzz.token_sort_ratio(item["name"], clf["name"])
                if sim < threshold:
                    r["status"] = "ОШИБКА"
                    errors.append("Наименование не соответствует")
                    r["expected"] = clf["name_orig"]
                    r["recommendation"] = f'Ожидается: «{clf["name_orig"]}» (схожесть {sim}%)'

                if clf["unit"] and item["unit"]:
                    if item["unit"] != clf["unit"]:
                        r["status"] = "ОШИБКА"
                        errors.append("Ед. изм. не соответствует")
                        unit_note = clf["unit_orig"]
                        if r["expected"]:
                            r["expected"] += f"; ед. изм.: {unit_note}"
                        else:
                            r["expected"] = unit_note
                        sep = "; " if r["recommendation"] else ""
                        r["recommendation"] += f"{sep}Ожидаемая ед. изм.: «{unit_note}»"

        r["error_type"] = " / ".join(errors)
        results.append(r)

    return results


# ── Формирование Excel-отчёта ──────────────────────────────────────────────────

def write_report(results: list, output: str):
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Ошибки"
    write_header(ws1, [
        "№ строки", "Кодификатор", "Наименование (смета)",
        "Ед. изм. (смета)", "Тип ошибки", "Ожидаемое значение", "Рекомендация",
    ])
    for r in results:
        if r["status"] == "ОШИБКА":
            ws1.append([
                r["row_num"], r["code"], r["name_est"],
                r["unit_est"], r["error_type"], r["expected"], r["recommendation"],
            ])
            for cell in ws1[ws1.max_row]:
                cell.fill = FILL_ERR
    ws1.freeze_panes = "A2"
    auto_width(ws1)

    ws2 = wb.create_sheet("Статистика")
    total  = len(results)
    ok_n   = sum(1 for r in results if r["status"] == "ОК")
    err_n  = sum(1 for r in results if r["status"] == "ОШИБКА")
    warn_n = sum(1 for r in results if r["status"] == "УТОЧНЕНИЕ")

    by_type: dict = {}
    for r in results:
        for t in r["error_type"].split(" / "):
            if t:
                by_type[t] = by_type.get(t, 0) + 1

    write_header(ws2, ["Показатель", "Значение"])
    stat_rows = [
        ("Всего строк проверено",       total),
        ("Корректных строк (ОК)",         ok_n),
        ("Строк с ошибками",              err_n),
        ("Требуют уточнения детализации", warn_n),
        ("", ""),
        ("Типы ошибок / предупреждений", ""),
    ]
    for t, cnt in by_type.items():
        stat_rows.append((f"  — {t}", cnt))
    stat_rows += [
        ("", ""),
        ("Процент корректных строк",   f"{ok_n  / total * 100:.1f}%" if total else "—"),
        ("Процент строк с ошибками",  f"{err_n / total * 100:.1f}%" if total else "—"),
    ]
    for row in stat_rows:
        ws2.append(list(row))
    ws2.column_dimensions["A"].width = 44
    ws2.column_dimensions["B"].width = 22

    ws3 = wb.create_sheet("Детализация")
    write_header(ws3, [
        "№ строки", "Кодификатор", "Наименование (смета)",
        "Ед. изм. (смета)", "Статус", "Тип ошибки / Примечание",
    ])
    fill_map = {"ОК": FILL_OK, "УТОЧНЕНИЕ": FILL_WARN, "ОШИБКА": FILL_ERR}
    for r in results:
        ws3.append([
            r["row_num"], r["code"], r["name_est"],
            r["unit_est"], r["status"], r["error_type"],
        ])
        for cell in ws3[ws3.max_row]:
            cell.fill = fill_map.get(r["status"], FILL_OK)
    ws3.freeze_panes = "A2"
    auto_width(ws3)

    wb.save(output)


# ── Точка входа ────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Сверка сметы генподрядчика с классификатором СМР",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Примеры:
  python checker.py classifier.xlsx estimate.xlsx
  python checker.py classifier.xlsx estimate.xlsx -o my_report.xlsx -t 80
        """,
    )
    parser.add_argument("classifier", help="Путь к файлу классификатора (.xlsx)")
    parser.add_argument("estimate",   help="Путь к файлу сметы (.xlsx)")
    parser.add_argument("-o", "--output",    default="report.xlsx",
                        help="Имя файла отчёта (default: report.xlsx)")
    parser.add_argument("-t", "--threshold", default=85, type=int,
                        help="Порог нечёткого сравнения 0-100 (default: 85)")
    args = parser.parse_args()

    print(f"[1/4] Загрузка классификатора : {args.classifier}")
    clf = load_classifier(args.classifier)
    print(f"      Позиций загружено        : {len(clf)}")

    print(f"[2/4] Загрузка сметы          : {args.estimate}")
    est = load_estimate(args.estimate)
    print(f"      Строк загружено          : {len(est)}")

    if not est:
        print("[ERROR] Строки сметы не найдены. Проверьте формат файла.")
        sys.exit(1)

    print(f"[3/4] Проверка (порог: {args.threshold}%)...")
    results = check(clf, est, args.threshold)

    print(f"[4/4] Формирование отчёта     : {args.output}")
    write_report(results, args.output)

    total  = len(results)
    ok_n   = sum(1 for r in results if r["status"] == "ОК")
    err_n  = sum(1 for r in results if r["status"] == "ОШИБКА")
    warn_n = sum(1 for r in results if r["status"] == "УТОЧНЕНИЕ")

    print()
    print("=" * 54)
    print("  ИТОГИ ПРОВЕРКИ")
    print("=" * 54)
    print(f"  Всего строк            : {total}")
    if total:
        print(f"  Корректных (ОК)        : {ok_n}  ({ok_n  / total * 100:.1f}%)")
        print(f"  Ошибок                 : {err_n}  ({err_n / total * 100:.1f}%)")
    print(f"  Требуют уточнения      : {warn_n}")

    if err_n or warn_n:
        print()
        print("  Типы проблем:")
        by_type: dict = {}
        for r in results:
            for t in r["error_type"].split(" / "):
                if t:
                    by_type[t] = by_type.get(t, 0) + 1
        for t, cnt in by_type.items():
            print(f"    - {t}: {cnt}")

    print("=" * 54)
    print(f"  Отчёт сохранён: {args.output}")
    print()


if __name__ == "__main__":
    main()
